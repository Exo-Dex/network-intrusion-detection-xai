"""
Compiles the IEEE conference draft markdown into a styled Microsoft Word (.docx) manuscript.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_manuscript(md_path, out_path):
    if not os.path.exists(md_path):
        print(f"Error: {md_path} not found.")
        return
        
    doc = docx.Document()
    
    # Page setup
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        if line.startswith('# '):
            p = doc.add_paragraph()
            run = p.add_run(line[2:])
            run.font.size = Pt(18)
            run.font.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            p = doc.add_paragraph()
            run = p.add_run(line[3:])
            run.font.size = Pt(12)
            run.font.bold = True
        elif line.startswith('### '):
            p = doc.add_paragraph()
            run = p.add_run(line[4:])
            run.font.size = Pt(10.5)
            run.font.bold = True
        elif line.startswith('|') and '---' not in line:
            # Table row
            doc.add_paragraph(line)
        else:
            doc.add_paragraph(line)
            
    doc.save(out_path)
    print(f"Compiled manuscript saved to: {out_path}")

if __name__ == '__main__':
    src = '/working_dir/c_2cc455fc38c695ba/network-intrusion-detection-xai/docs/research_paper/ieee_paper_draft.md'
    dst = '/working_dir/c_2cc455fc38c695ba/network-intrusion-detection-xai/docs/research_paper/nids_xai_ieee_compiled.docx'
    build_manuscript(src, dst)
